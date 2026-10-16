import marimo

__generated_with = "0.24.2"
app = marimo.App(width="medium")


@app.cell
def _():
    import marimo as mo
    import pandas as pd

    return mo, pd


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    # Pandas Input/Output

    ## Objectives

    - Read and write CSV and JSON files with Pandas.
    - Tell `read_csv` when a file is separated by semicolons.
    - Read an ARFF file and an Excel workbook.
    - Flatten a nested JSON object with `json_normalize`.

    ## Background

    Pandas can load a file into a DataFrame and write a DataFrame back out. The cells that use a web address need a network connection. Saving `auto.csv` and `countries.json` writes those files in this folder.

    ## Datasets Used

    - **Automobile Dataset** from the UCI Machine Learning Repository, for CSV files.
    - **Higher Education Students Performance Evaluation Dataset**, a CSV file separated by semicolons.
    - **Student Academics Performance Dataset**, an ARFF file.
    - **Countries**, a JSON file.
    - **Immunotherapy Dataset**, an Excel file.
    """)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Some dataset repositories

    **UCI Machine Learning Repository**

    The UCI Machine Learning Repository, [https://archive.ics.uci.edu/ml/index.php](https://archive.ics.uci.edu/ml/index.php), is a collection of datasets used by the machine learning community. It was created as an ftp archive in 1987 by David Aha and fellow graduate students at UC Irvine.

    For each dataset, the repository records the source, a description of the data, the attributes, and relevant papers.

    **Data.World**

    Data.world, [https://data.world/](https://data.world/), is a catalog of datasets in many formats.
    """)
    return


@app.cell
def _(pd):
    # Controlling the number of columns a DataFrame shows
    pd.set_option('display.max_columns', 8)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Reading data from a `.csv` file

    Comma-separated values (CSV) is one of the most common ways to store and share a table.

    `read_csv` reads a CSV file and returns a DataFrame. This file has no header row, so the column names are passed with `names`. The character `?` means missing, so `na_values='?'` turns it into `NaN`.

    The Automobile dataset is [https://archive.ics.uci.edu/ml/datasets/automobile](https://archive.ics.uci.edu/ml/datasets/automobile).
    """)
    return


@app.cell
def _():
    headers = [
        'symboling', 'normalized_losses', 'make', 'fuel_type', 'aspiration',
        'num_doors', 'body_style', 'drive_wheels', 'engine_location',
        'wheel_base', 'length', 'width', 'height', 'curb_weight',
        'engine_type', 'num_cylinders', 'engine_size', 'fuel_system',
        'bore', 'stroke', 'compression_ratio', 'horsepower', 'peak_rpm',
        'city_mpg', 'highway_mpg', 'price',
    ]
    return (headers,)


@app.cell
def _(headers, mo, pd):
    df_auto = pd.read_csv(
        'https://archive.ics.uci.edu/ml/machine-learning-databases/autos/imports-85.data',
        header=None,
        names=headers,
        na_values='?',
    )
    mo.Html(df_auto.head().to_html())
    return (df_auto,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    `shape` is `(rows, columns)`. The automobile table has 205 rows and 26 columns.
    """)
    return


@app.cell
def _(df_auto):
    df_auto.shape
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    `to_csv` writes the DataFrame to a file. This first save includes the index, so reading that file adds an extra column, `Unnamed: 0`. That column is the old index.
    """)
    return


@app.cell
def _(df_auto, mo, pd):
    df_auto.to_csv('auto.csv')
    with_index = pd.read_csv('auto.csv')
    mo.Html(with_index.head().to_html())
    return (with_index,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The extra column makes 27 columns instead of 26.
    """)
    return


@app.cell
def _(with_index):
    with_index.shape
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The second save passes `index=False`, and the following read does not have that extra column.
    """)
    return


@app.cell
def _(df_auto, mo, pd, with_index):
    with_index  # this save must run after the file with the index was read
    df_auto.to_csv('auto.csv', index=False)
    df2_auto = pd.read_csv('auto.csv')
    mo.Html(df2_auto.head().to_html())
    return (df2_auto,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The table has 26 columns again, the same shape as `df_auto`.
    """)
    return


@app.cell
def _(df2_auto):
    df2_auto.shape
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    `describe` summarizes the numeric columns of the automobile data.
    """)
    return


@app.cell
def _(df_auto, mo):
    mo.Html(df_auto.describe().to_html())
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    A CSV file is not always separated by commas. It may use semicolons, tabs, or another character. Set that character with `sep`.

    The Higher Education Students Performance Evaluation dataset is [https://archive.ics.uci.edu/ml/datasets/Higher+Education+Students+Performance+Evaluation+Dataset](https://archive.ics.uci.edu/ml/datasets/Higher+Education+Students+Performance+Evaluation+Dataset).

    The first read uses the default comma. The whole line lands in one column.
    """)
    return


@app.cell
def _(pd):
    df_hst_comma = pd.read_csv(
        'https://archive.ics.uci.edu/ml/machine-learning-databases/00623/DATA.csv',
        na_values='?',
    )
    df_hst_comma.head()
    return (df_hst_comma,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The shape confirms it: there is only one column.
    """)
    return


@app.cell
def _(df_hst_comma):
    df_hst_comma.shape
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The second read sets `sep=';'`. Each value now has its own column.
    """)
    return


@app.cell
def _(mo, pd):
    df_hst = pd.read_csv(
        'https://archive.ics.uci.edu/ml/machine-learning-databases/00623/DATA.csv',
        sep=';',
        na_values='?',
    )
    mo.Html(df_hst.head().to_html())
    return (df_hst,)


@app.cell
def _(df_hst):
    df_hst.shape
    return


@app.cell
def _(df_hst, mo):
    mo.Html(df_hst.describe().to_html())
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ### Try it yourself

    The first cell reads the local `auto.csv` file and shows its first three rows. The second cell shows its shape. Change the number of rows in `head` and run the first cell again.
    """)
    return


@app.cell
def _(df2_auto, mo, pd):
    df2_auto  # read the file after it was saved without the index
    check = pd.read_csv('auto.csv')
    mo.Html(check.head(3).to_html())
    return (check,)


@app.cell
def _(check):
    check.shape
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ### Try it yourself

    Read the text `"a;b\n1;?\n3;4\n"` with `read_csv`. The separator is a semicolon, and `?` means missing. `StringIO` lets `read_csv` read a string as if it were a file. Replace `None` with your code. What does the `?` become?
    """)
    return


@app.cell
def _():
    student_semicolon = None
    student_semicolon
    return (student_semicolon,)


@app.cell(hide_code=True)
def _(mo):
    mo.accordion({"Show answers": mo.md(r'''
```python
from io import StringIO
student_semicolon = pd.read_csv(StringIO("a;b\n1;?\n3;4\n"), sep=';', na_values='?')
```
`NaN`. `sep` tells Pandas where one column ends, and `na_values` names the text that should be missing data.
''')}, lazy=True)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Reading data from an `.arff` file

    An ARFF file (Attribute-Relation File Format) is a text file from the Machine Learning Project at the University of Waikato. It has a header section and a data section.

    `scipy.io.arff.loadarff` reads both sections. The Student Academics file is downloaded from [https://archive.ics.uci.edu/ml/machine-learning-databases/00467/Sapfile1.arff](https://archive.ics.uci.edu/ml/machine-learning-databases/00467/Sapfile1.arff). The download uses `urlopen` from Python's standard library. Text values arrive as byte strings, so each text column is decoded to ordinary text.
    """)
    return


@app.cell
def _():
    # For reading an arff file, we need to import
    from io import StringIO
    from urllib.request import urlopen
    from scipy.io.arff import loadarff

    return StringIO, loadarff, urlopen


@app.cell
def _(urlopen):
    # Fetch the content from the URL
    stAcademic_url = 'https://archive.ics.uci.edu/ml/machine-learning-databases/00467/Sapfile1.arff'
    with urlopen(stAcademic_url) as downloaded:
        status_code = downloaded.status
        payload = downloaded.read()
    return payload, status_code, stAcademic_url


@app.cell
def _(StringIO, loadarff, payload, pd, status_code):
    # Check if the request was successful
    if status_code == 200:
        # Parse the ARFF data
        arff_content = payload.decode('utf-8')
        arff_data = StringIO(arff_content)

        # Get the data and metadata
        arff_rows, meta = loadarff(arff_data)

        # Convert to Pandas DataFrame
        df = pd.DataFrame(arff_rows)

        # Decode byte strings if necessary
        for column in df.select_dtypes([object]).columns:
            if df[column].apply(lambda x: isinstance(x, bytes)).any():
                df[column] = df[column].str.decode('utf-8')
    else:
        df = pd.DataFrame()
        meta = None
        print(f'Failed to download the file: {status_code}')
    return df, meta


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The data section becomes the DataFrame:
    """)
    return


@app.cell
def _(df, mo):
    mo.Html(df.head().to_html())
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The header section becomes the metadata. It lists each attribute with its type, and the allowed values of each nominal attribute.
    """)
    return


@app.cell
def _(meta):
    print(meta)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The metadata knows the column names:
    """)
    return


@app.cell
def _(meta):
    column_names = meta.names()
    column_names
    return (column_names,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    `describe` summarizes the numeric columns.
    """)
    return


@app.cell
def _(df, mo):
    mo.Html(df.describe().to_html())
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Reading data from a `.json` file

    JavaScript Object Notation (JSON) is a common way to share data on the web.

    `read_json` reads a JSON file and returns a DataFrame. This example uses a countries file from data.world.
    """)
    return


@app.cell
def _(mo, pd):
    df_countries = pd.read_json('https://query.data.world/s/6wc2blqdaxd6s2k3xmzmfx3zb72v7l')
    mo.Html(df_countries.head().to_html())
    return (df_countries,)


@app.cell
def _(df_countries):
    df_countries.shape
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    `describe` summarizes the numeric columns.
    """)
    return


@app.cell
def _(df_countries, mo):
    mo.Html(df_countries.describe().to_html())
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    `to_json` writes that DataFrame to `countries.json` in this folder.
    """)
    return


@app.cell
def _(df_countries):
    df_countries.to_json('countries.json')
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    JSON can be nested. `json_normalize` flattens it into a table.

    In this example each person has a name, and each car also has a name. `meta_prefix` and `record_prefix` keep those names apart. The result has the columns `person_name`, `person_age`, `car_name`, and `car_models`.
    """)
    return


@app.cell
def _():
    from pandas import json_normalize

    # a simple JSON object
    data = [
        {
            'name': 'John',
            'age': 30,
            'cars': [
                {'name': 'Ford', 'models': ['Fiesta', 'Focus', 'Mustang']},
                {'name': 'BMW', 'models': ['320', 'X3', 'X5']},
            ],
        },
        {
            'name': 'Peter',
            'age': 46,
            'cars': [
                {'name': 'Mercedes', 'models': ['E Class', 'S Class']},
            ],
        },
    ]

    # use json_normalize to flatten the data
    # give the fact that name is a column in the data, we need to specify the meta_prefix
    # and record_prefix to avoid conflicts
    # The final dataframe will contain columns: person_name, person_age, car_name, car_models
    cars = json_normalize(
        data,
        record_path='cars',
        meta=['name', 'age'],
        meta_prefix='person_',
        record_prefix='car_',
    )
    cars
    return json_normalize, data, cars


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ### Try it yourself

    Flatten `data` the same way the lesson does: the records are under `'cars'`, and the person fields are `name` and `age`. Use the prefixes `person_` and `car_`. Replace `None` with your code. Why are the prefixes needed?
    """)
    return


@app.cell
def _():
    student_flat = None
    student_flat
    return (student_flat,)


@app.cell(hide_code=True)
def _(mo):
    mo.accordion({"Show answers": mo.md(r'''
```python
student_flat = json_normalize(
    data,
    record_path='cars',
    meta=['name', 'age'],
    meta_prefix='person_',
    record_prefix='car_',
)
```
Both the person and the car have a field named `name`. The prefixes keep those two columns from being given the same name.
''')}, lazy=True)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Reading data from a `.xlsx` file

    `read_excel` reads an Excel workbook and returns a DataFrame. This cell loads the Immunotherapy dataset, [https://archive.ics.uci.edu/ml/datasets/Immunotherapy+Dataset](https://archive.ics.uci.edu/ml/datasets/Immunotherapy+Dataset). `openpyxl` is the library Pandas uses to read `.xlsx` files.
    """)
    return


@app.cell
def _(mo, pd):
    import openpyxl

    df_immuno = pd.read_excel(
        'https://archive.ics.uci.edu/ml/machine-learning-databases/00428/Immunotherapy.xlsx',
        engine='openpyxl',
    )
    mo.Html(df_immuno.head().to_html())
    return df_immuno, openpyxl


@app.cell
def _(df_immuno):
    df_immuno.shape
    return


@app.cell
def _(df_immuno, mo):
    mo.Html(df_immuno.describe().to_html())
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ### Try it yourself

    Build a one-column DataFrame from `{'City': ['Rome', 'Oslo']}` and convert it to CSV text with `index=False`. Replace `None` with your code. Do the row numbers appear in that text?
    """)
    return


@app.cell
def _():
    student_csv_text = None
    print(student_csv_text)
    return (student_csv_text,)


@app.cell(hide_code=True)
def _(mo):
    mo.accordion({"Show answers": mo.md(r'''
```python
student_cities = pd.DataFrame({'City': ['Rome', 'Oslo']})
student_csv_text = student_cities.to_csv(index=False)
```
No. `index=False` leaves the index out of the file. Without it, the row numbers are written as an extra column.
''')}, lazy=True)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Conclusions

    **Key takeaways:**

    - `read_csv` and `to_csv` move a table to and from a CSV file. `index=False` leaves the index out of the file.
    - `sep` sets the character between columns. `na_values` names the text that should become missing data. `names` supplies the column headers.
    - An ARFF file can be loaded with `loadarff` and turned into a DataFrame. Byte strings need to be decoded.
    - `read_json` and `to_json` handle JSON tables. `json_normalize` flattens nested records, and prefixes keep duplicate field names apart.
    - `read_excel` reads an `.xlsx` file. `describe` is a quick summary of the numeric columns after the load.
    """)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Check your understanding

    1. What does `names` supply to `read_csv`?
    2. What does `sep=';'` tell `read_csv`?
    3. What does `na_values='?'` do?
    4. What does `index=False` leave out of a CSV file?
    5. What two parts does an ARFF file contain, and what has to be done to its text values?
    6. What does `json_normalize` do to nested records?
    7. Why does the cars example use `meta_prefix` and `record_prefix`?
    """)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.accordion({"Show answers": mo.md(r"""
1. **The column headers:** The automobile file has no header row, so the names are passed in separately.
2. **The separator:** A semicolon, rather than a comma, stands between the columns.
3. **It marks missing data:** The character `?` becomes `NaN`.
4. **The index:** The row labels are not written as an extra column.
5. **A header and the data:** `loadarff` reads both. Text values arrive as byte strings and are decoded to ordinary text.
6. **It flattens them:** Nested fields become columns of one table.
7. **Both levels have a `name`:** The prefixes keep the person's name and the car's name in different columns.
""")}, lazy=True)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## References

    - VanderPlas, J. (2017). *Python Data Science Handbook: Essential Tools for Working with Data*. O'Reilly Media. Chapter 3.
    """)
    return


if __name__ == "__main__":
    app.run()
